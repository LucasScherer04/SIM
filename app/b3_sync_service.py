import io
import re
import csv
import json
import base64
import zipfile
import unicodedata
import xml.etree.ElementTree as ET
from datetime import datetime, date
from typing import Dict, List, Any, Tuple, Optional

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

from .database import get_conn, migrate_fixed_income_columns
from .market_service import lookup_ticker_live, fetch_macro_indicators

KNOWN_ETFS_BR = {
    "BOVA11", "IVVB11", "SMAL11", "HASH11", "DIVO11", "WRLD11", "GOLD11",
    "NASD11", "SPXI11", "XFIX11", "BOVV11", "BRAX11", "ECOO11", "SMAC11",
    "PIBB11", "GOVE11", "MATB11", "FIND11", "ISUS11", "SHOT11", "ACWI11",
    "EURP11", "ASIA11", "XINA11", "QBTC11", "ETHE11", "BITH11", "DEFI11",
    "WEB311", "LFTS11", "IMAB11", "IB5M11", "B5P211", "NTNS11", "NDIV11",
}

KNOWN_STOCK_UNITS_BR = {
    "TAEE11", "SANB11", "KLBN11", "ALUP11", "SAPR11", "ENGI11", "BPAC11",
    "IGTI11", "PPLA11", "MODL11", "RNEW11", "AESB11",
}

# Classes de ativos estritamente custodiadas na B3 (mercado de capitais brasileiro)
B3_CUSTODY_TRANSACTION_CLASSES = {"ACAO_BR", "FII", "ETF_BR"}
B3_CUSTODY_FIXED_INCOME_CLASSES = {"TESOURO", "RENDA_FIXA"}

# Classes que NUNCA são custodiadas na B3 e JAMAIS devem ser apagadas ou alteradas pela sincronização B3
NON_B3_CLASSES = {"STOCK_US", "CRIPTO", "CAIXINHA", "FUNDO_INV"}


def ensure_b3_sync_tables():
    conn = get_conn()
    conn.execute("""
        CREATE TABLE IF NOT EXISTS b3_sync_logs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL DEFAULT 1,
            synced_at TEXT NOT NULL,
            filename TEXT NOT NULL,
            sync_mode TEXT NOT NULL,
            report_types TEXT DEFAULT '',
            positions_synced INTEGER DEFAULT 0,
            transactions_synced INTEGER DEFAULT 0,
            fixed_income_synced INTEGER DEFAULT 0,
            dividends_synced INTEGER DEFAULT 0,
            skipped_duplicates INTEGER DEFAULT 0,
            summary_text TEXT DEFAULT ''
        )
    """)
    cols = {r[1] for r in conn.execute("PRAGMA table_info(b3_sync_logs)").fetchall()}
    if "user_id" not in cols:
        conn.execute("ALTER TABLE b3_sync_logs ADD COLUMN user_id INTEGER DEFAULT 1")
    conn.commit()
    conn.close()
    migrate_fixed_income_columns()


def get_current_portfolio_holdings(cur_or_conn, user_id: int = 1) -> Dict[str, Dict[str, Any]]:
    """
    Computes current held quantity and weighted average purchase price (Preço Médio)
    for every asset in the transactions table.
    Identifies whether each position has real historical trades or only synthetic B3 imports.
    """
    rows = cur_or_conn.execute("SELECT * FROM transactions WHERE user_id = ? ORDER BY tx_date ASC, id ASC", (user_id,)).fetchall()
    holdings: Dict[str, Dict[str, Any]] = {}
    for r in rows:
        t = str(r["ticker"] or "").upper().strip()
        if not t:
            continue
        if t not in holdings:
            holdings[t] = {
                "ticker": t,
                "quantity": 0.0,
                "avg_price": 0.0,
                "asset_class": r["asset_class"],
                "broker": r["broker"] or "",
                "has_synthetic_pos": False,
                "synthetic_tx_ids": [],
            }
        item = holdings[t]
        notes_raw = r["notes"] or ""
        notes = "".join(ch for ch in unicodedata.normalize("NFD", notes_raw) if unicodedata.category(ch) != "Mn").lower()
        if any(k in notes for k in ("sincronizacao posicao b3", "posicao b3", "ajuste de quantidade b3", "ajuste de custodia b3", "ajuste b3", "preco medio informado")):
            item["has_synthetic_pos"] = True
            item["synthetic_tx_ids"].append(r["id"])

        q = float(r["quantity"])
        p = float(r["unit_price"])
        f = float(r["fees"] or 0.0)
        ttype = str(r["tx_type"] or "").upper().strip()
        if ttype in ("BUY", "COMPRA"):
            prev_cost = item["quantity"] * item["avg_price"]
            new_cost = prev_cost + (q * p) + f
            item["quantity"] += q
            item["avg_price"] = new_cost / item["quantity"] if item["quantity"] > 0 else 0.0
            item["broker"] = r["broker"] or item["broker"]
            if r["asset_class"]:
                item["asset_class"] = r["asset_class"]
        elif ttype in ("SELL", "VENDA"):
            item["quantity"] = max(0.0, item["quantity"] - q)
            if item["quantity"] <= 1e-8:
                item["avg_price"] = 0.0
    return holdings


def _norm_text(val: Any) -> str:
    """Normalizes string removing accents and lowercasing for flexible header/value matching."""
    if val is None:
        return ""
    s = str(val).strip()
    s = unicodedata.normalize("NFD", s)
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn")
    return s.lower().strip()


def _parse_br_float(val: Any) -> float:
    """Parses Brazilian or international number strings ('R$ 1.234,56', '1.000', '1234.56', '-', etc.) into float."""
    if val is None:
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip()
    if not s or s in ("-", "--", "N/A", "n/a", "null", "None"):
        return 0.0
    # Strip currency symbols, percent, spaces, non-breaking spaces
    s = re.sub(r"[R\$US\$\s%\xa0]", "", s)
    if not s:
        return 0.0
    # Check if both '.' and ',' exist
    if "," in s and "." in s:
        if s.rfind(",") > s.rfind("."):
            # Brazilian format: 1.234,56 -> 1234.56
            s = s.replace(".", "").replace(",", ".")
        else:
            # International format: 1,234.56 -> 1234.56
            s = s.replace(",", "")
    elif "," in s:
        # Brazilian decimal: 1234,56 -> 1234.56
        s = s.replace(",", ".")
    elif "." in s:
        # Check if dot is Brazilian thousand separator: e.g. '1.000', '25.000', '1.000.000'
        # In Brazilian spreadsheets and CSVs, dot alone is a thousand separator when:
        # 1) It matches standard thousand separator groupings e.g. '1.000', '25.000', '1.000.000' (not starting with 0)
        # 2) Or if it has multiple dots e.g. 1.000.000
        if re.match(r"^-?[1-9]\d{0,2}(?:\.\d{3})+$", s):
            s = s.replace(".", "")
        elif s.count(".") > 1:
            s = s.replace(".", "")
    try:
        return float(s)
    except Exception:
        m = re.search(r"[-+]?\d+(?:\.\d+)?", s)
        return float(m.group(0)) if m else 0.0


def _parse_br_date(val: Any) -> str:
    """Parses dates from Excel (datetime/date/serial or 'DD/MM/YYYY' or 'YYYY-MM-DD') into 'YYYY-MM-DD'."""
    if val is None:
        return date.today().strftime("%Y-%m-%d")
    if isinstance(val, datetime):
        return val.strftime("%Y-%m-%d")
    if isinstance(val, date):
        return val.strftime("%Y-%m-%d")
    s = str(val).strip()
    if not s or s in ("-", "--", "Sem vencimento"):
        return date.today().strftime("%Y-%m-%d")

    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y", "%Y/%m/%d"):
        try:
            return datetime.strptime(s[:10].strip(), fmt).strftime("%Y-%m-%d")
        except Exception:
            continue

    m_br = re.search(r"(\d{2})/(\d{2})/(\d{4})", s)
    if m_br:
        d, m, y = m_br.groups()
        return f"{y}-{m}-{d}"
    m_iso = re.search(r"(\d{4})-(\d{2})-(\d{2})", s)
    if m_iso:
        y, m, d = m_iso.groups()
        return f"{y}-{m}-{d}"

    return date.today().strftime("%Y-%m-%d")


IGNORED_TICKER_WORDS = {
    "TOTAL", "SUBTOTAL", "SALDO", "VALOR", "DISPONIVEL", "CREDITO", "DEBITO",
    "DATA", "RESUMO", "CUSTODIA", "POSICAO", "ATIVO", "EXTRATO", "CONTA", "MOVIMENTACAO",
    "TITULO", "APLICACAO", "RESGATE", "RENDIMENTO", "IRRF", "IOF", "TAXA", "B3",
    "COMPRA", "VENDA", "LIQUIDACAO", "PROVENTOS", "DIVIDENDO", "JSCP", "JUROS"
}


def _extract_ticker_and_name(product_str: str, code_str: str = "") -> Tuple[str, str]:
    """
    Extracts normalized B3 ticker (e.g., PETR4F -> PETR4) and asset name from B3 'Produto' and 'Código de Negociação'.
    Examples:
      product_str='PETR4 - PETROLEO BRASILEIRO S.A. PETROBRAS', code_str='PETR4F' -> ('PETR4', 'PETROLEO BRASILEIRO S.A. PETROBRAS')
      product_str='MXRF11 - FII MAXI RENDA', code_str='' -> ('MXRF11', 'FII MAXI RENDA')
    """
    raw_code = str(code_str or "").strip().upper().replace(".SA", "")
    raw_prod = str(product_str or "").strip()

    extracted_name = raw_prod
    if " - " in raw_prod:
        parts = raw_prod.split(" - ", 1)
        left_candidate = parts[0].strip().upper()
        if not raw_code and re.match(r"^[A-Z]{4}\d{1,2}[FB]?$", left_candidate):
            raw_code = left_candidate
            extracted_name = parts[1].strip()
        elif left_candidate == raw_code or left_candidate == raw_code.rstrip("F"):
            extracted_name = parts[1].strip()

    if not raw_code:
        # Search for B3 ticker pattern inside product_str
        m = re.search(r"\b([A-Z]{4}\d{1,2}[F]?)\b", raw_prod.upper())
        if m:
            raw_code = m.group(1)
        else:
            m2 = re.match(r"^([A-Z]{1,5})\b", raw_prod.upper())
            if m2:
                candidate = m2.group(1)
                if candidate not in IGNORED_TICKER_WORDS:
                    raw_code = candidate

    # Strip fractional 'F' suffix from B3 stocks/FIIs/ETFs (e.g., PETR4F -> PETR4, WEGE3F -> WEGE3, MXRF11F -> MXRF11)
    if re.match(r"^[A-Z]{4}\d{1,2}F$", raw_code):
        raw_code = raw_code[:-1]

    if raw_code in IGNORED_TICKER_WORDS:
        raw_code = ""

    return raw_code, (extracted_name or raw_code)


def _infer_asset_class(ticker: str, sheet_name: str = "", product_name: str = "", tipo_str: str = "") -> Tuple[str, str]:
    """
    Infers (asset_class, currency) for an asset found in a B3 report.
    On B3 (Brasil, Bolsa, Balcão), assets are strictly Brazilian custody:
    ACAO_BR, FII, or ETF_BR, and the settlement currency is always BRL.
    Assets outside B3 (STOCK_US, CRIPTO) are NEVER inferred from a B3 report.
    """
    t = ticker.upper().strip()
    conn = get_conn()
    row = conn.execute("SELECT asset_class, currency FROM assets_catalog WHERE ticker = ?", (t,)).fetchone()
    conn.close()
    if row and row["asset_class"] in B3_CUSTODY_TRANSACTION_CLASSES:
        return row["asset_class"], "BRL"

    sh_norm = _norm_text(sheet_name)
    prod_norm = _norm_text(product_name)
    tipo_norm = _norm_text(tipo_str)

    if t in KNOWN_ETFS_BR or "etf" in sh_norm or "etf" in prod_norm or "etf" in tipo_norm or "ishares" in prod_norm or "it now" in prod_norm or "trend " in prod_norm:
        return "ETF_BR", "BRL"

    if t in KNOWN_STOCK_UNITS_BR:
        return "ACAO_BR", "BRL"

    if "fii" in sh_norm or "fundo" in sh_norm or "fii" in prod_norm or "fiagro" in prod_norm or "imobiliario" in prod_norm or "cota" in tipo_norm:
        if t.endswith("11") or t.endswith("11B"):
            return "FII", "BRL"

    if t.endswith("11") or t.endswith("11B"):
        # On B3, 11-ending tickers that aren't known units or ETFs are overwhelmingly FIIs / Fiagros
        return "FII", "BRL"

    if re.match(r"^[A-Z]{4}(3|4|5|6|7|8)$", t):
        return "ACAO_BR", "BRL"

    if re.match(r"^[A-Z]{4}(32|33|34|35|39)$", t) or "bdr" in sh_norm or "bdr" in tipo_norm:
        # BDRs traded on B3 in BRL
        return "ACAO_BR", "BRL"

    return "ACAO_BR", "BRL"


def _clean_broker_name(raw_broker: str) -> str:
    if not raw_broker:
        return "B3 / Corretora"
    b = str(raw_broker).strip()
    b_up = b.upper()
    if "XP INVESTIMENTOS" in b_up:
        return "XP Investimentos"
    if "RICO" in b_up:
        return "Rico"
    if "CLEAR" in b_up:
        return "Clear"
    if "NU INVEST" in b_up or "NUBANK" in b_up or "NU PAGAMENTOS" in b_up:
        return "Nubank / NuInvest"
    if "BTG" in b_up:
        return "BTG Pactual"
    if "INTER" in b_up:
        return "Banco Inter"
    if "ITAU" in b_up or "ITAÚ" in b_up:
        return "Itaú Corretora"
    if "BRADESCO" in b_up or "AGORA" in b_up or "ÁGORA" in b_up:
        return "Ágora / Bradesco"
    if "SANTANDER" in b_up or "TORO" in b_up:
        return "Toro / Santander"
    if "GENIAL" in b_up:
        return "Genial Investimentos"
    if "ORAMA" in b_up or "ÓRAMA" in b_up:
        return "Órama"
    if "BANCO DO BRASIL" in b_up or "BB " in b_up:
        return "BB Investimentos"
    if "C6" in b_up:
        return "C6 Bank"
    return b[:36].strip()


def _read_xlsx_fallback_stdlib(file_bytes: bytes) -> Dict[str, List[List[Any]]]:
    """
    Pure-Python stdlib fallback reader for .xlsx files using zipfile + xml.etree.ElementTree,
    used only if openpyxl is not installed on the system.
    """
    sheets: Dict[str, List[List[Any]]] = {}
    with zipfile.ZipFile(io.BytesIO(file_bytes)) as zf:
        shared_strings: List[str] = []
        if "xl/sharedStrings.xml" in zf.namelist():
            root = ET.fromstring(zf.read("xl/sharedStrings.xml"))
            for si in root.iter():
                if si.tag.endswith("}si") or si.tag == "si":
                    texts = [t.text or "" for t in si.iter() if t.tag.endswith("}t") or t.tag == "t"]
                    shared_strings.append("".join(texts))

        wb_root = ET.fromstring(zf.read("xl/workbook.xml"))
        sheet_names = []
        for sh in wb_root.iter():
            if sh.tag.endswith("}sheet") or sh.tag == "sheet":
                sheet_names.append(sh.attrib.get("name", f"Sheet{len(sheet_names)+1}"))

        sheet_files = sorted(
            [f for f in zf.namelist() if re.match(r"^xl/worksheets/sheet\d+\.xml$", f)],
            key=lambda x: int(re.search(r"(\d+)", x).group(1)),
        )

        for idx, sf in enumerate(sheet_files):
            s_name = sheet_names[idx] if idx < len(sheet_names) else f"Sheet{idx+1}"
            s_root = ET.fromstring(zf.read(sf))
            rows_data: List[List[Any]] = []
            for row_el in s_root.iter():
                if row_el.tag.endswith("}row") or row_el.tag == "row":
                    row_vals = []
                    for c_el in row_el:
                        if not (c_el.tag.endswith("}c") or c_el.tag == "c"):
                            continue
                        c_type = c_el.attrib.get("t", "")
                        v_text = ""
                        for child in c_el:
                            if child.tag.endswith("}v") or child.tag == "v":
                                v_text = child.text or ""
                            elif child.tag.endswith("}is") or child.tag == "is":
                                v_text = "".join(t.text or "" for t in child.iter() if t.tag.endswith("}t") or t.tag == "t")
                        if c_type == "s" and v_text.isdigit():
                            s_idx = int(v_text)
                            val = shared_strings[s_idx] if 0 <= s_idx < len(shared_strings) else v_text
                        else:
                            val = v_text
                        row_vals.append(val)
                    if any(str(x).strip() for x in row_vals):
                        rows_data.append(row_vals)
            sheets[s_name] = rows_data
    return sheets


def load_workbook_sheets(file_bytes: bytes, filename: str) -> Dict[str, List[List[Any]]]:
    """
    Loads all sheets from an .xlsx or .csv file into {sheet_name: [[col1, col2, ...], ...]}.
    """
    fname_lower = (filename or "").lower()
    is_xlsx = fname_lower.endswith(".xlsx") or file_bytes[:4] == b"PK\x03\x04"

    if is_xlsx:
        if HAS_OPENPYXL:
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
            sheets: Dict[str, List[List[Any]]] = {}
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                rows = []
                for row in ws.iter_rows(values_only=True):
                    if row and any(cell is not None and str(cell).strip() != "" for cell in row):
                        rows.append(list(row))
                if rows:
                    sheets[sheet_name] = rows
            return sheets
        else:
            return _read_xlsx_fallback_stdlib(file_bytes)

    # Otherwise parse as CSV / TSV text
    text = ""
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            text = file_bytes.decode(enc)
            break
        except Exception:
            continue

    if not text.strip():
        return {}

    # Detect delimiter (; or , or \t)
    first_lines = "\n".join(text.splitlines()[:5])
    if first_lines.count(";") >= first_lines.count(","):
        delimiter = ";"
    elif "\t" in first_lines and first_lines.count("\t") > first_lines.count(","):
        delimiter = "\t"
    else:
        delimiter = ","

    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    rows = [r for r in reader if r and any(str(c).strip() for c in r)]
    return {"CSV": rows} if rows else {}


KNOWN_HEADER_COLS = {
    "produto", "instituicao", "conta", "codigo de negociacao", "codigo", "ticker",
    "ativo", "cnpj", "cnpj da empresa", "cnpj do fundo", "tipo", "tipo de ativo",
    "quantidade", "preco de fechamento", "preco unitario", "preco medio", "preco",
    "valor atualizado", "valor aplicado", "valor bruto", "valor liquido", "valor da operacao", "valor",
    "emissor", "indexador", "data de emissao", "vencimento", "data do negocio",
    "tipo de movimentacao", "movimentacao", "entrada/saida", "codigo isin", "isin",
    "valor atualizado mtm", "valor atualizado curva", "mercado", "prazo/vencimento"
}


def _is_table_header_row(row: List[Any]) -> bool:
    """
    Determines robustly whether a row is a table header row (column names)
    and not a data row that happens to contain keywords like 'Produto' or 'Valor'.
    """
    if not row or len(row) < 3:
        return False
    norm_cells = [_norm_text(c) for c in row if c is not None and str(c).strip()]
    if len(norm_cells) < 3:
        return False

    header_col_count = 0
    data_cell_count = 0

    for c in norm_cells:
        # Check if cell is a date: DD/MM/YYYY or YYYY-MM-DD
        if re.match(r"^\d{2}/\d{2}/\d{4}$", c) or re.match(r"^\d{4}-\d{2}-\d{2}$", c):
            data_cell_count += 1
            continue
        # Check if cell is a number or currency (e.g. 100, 16.79, 16,79, 1514.57)
        cleaned = re.sub(r"[r\$us\$\s%]", "", c)
        if cleaned and re.match(r"^-?\d+(?:[.,]\d+)?$", cleaned):
            data_cell_count += 1
            continue

        c_stripped = re.sub(r"[\(\)\[\]]", "", c).strip()
        if c_stripped in KNOWN_HEADER_COLS:
            header_col_count += 1
        elif any(c_stripped == k or c_stripped.startswith(k + " ") or c_stripped.endswith(" " + k) for k in KNOWN_HEADER_COLS):
            header_col_count += 1

    # Real header rows contain column titles, never data numbers/dates, and match multiple recognized headers
    if data_cell_count >= 1:
        return False
    return header_col_count >= 3


def _find_header_row(rows: List[List[Any]]) -> Tuple[int, List[str]]:
    """
    Scans the first 15 rows of a sheet to find the table header row (in case B3 adds title rows at the top).
    """
    best_idx = 0
    best_score = -1
    best_headers: List[str] = []

    for idx, row in enumerate(rows[:15]):
        if _is_table_header_row(row):
            norm_cells = [_norm_text(c) for c in row]
            return idx, norm_cells

        norm_cells = [_norm_text(c) for c in row]
        score = sum(1 for c in norm_cells if any(k == c or c.startswith(k + " ") for k in KNOWN_HEADER_COLS))
        if score > best_score:
            best_score = score
            best_idx = idx
            best_headers = norm_cells

    return best_idx, best_headers


def _col_idx(headers: List[str], candidates: List[str]) -> int:
    """Finds the first column index matching any candidate keyword."""
    for cand in candidates:
        for i, h in enumerate(headers):
            if cand == h:
                return i
    for cand in candidates:
        for i, h in enumerate(headers):
            if cand in h:
                return i
    return -1


def _get_cell(row: List[Any], idx: int, default: Any = "") -> Any:
    if idx < 0 or idx >= len(row):
        return default
    val = row[idx]
    return default if val is None else val


def _partition_sheet_into_sections(sheet_name: str, rows: List[List[Any]]) -> List[Tuple[str, List[List[Any]]]]:
    """
    If a sheet or CSV contains multiple stacked tables (e.g. Ações, FIIs, Renda Fixa, Tesouro Direto),
    partitions rows into individual sections [(section_title, section_rows), ...].
    Uses strict column header validation to prevent data rows with words like 'Produto' or 'Valor' from splitting tables.
    """
    if len(rows) < 2:
        return [(sheet_name, rows)]

    header_indices = []
    for idx, r in enumerate(rows):
        if _is_table_header_row(r):
            header_indices.append(idx)

    if len(header_indices) <= 1:
        return [(sheet_name, rows)]

    sections: List[Tuple[str, List[List[Any]]]] = []
    for i, h_idx in enumerate(header_indices):
        end_idx = header_indices[i + 1] if i + 1 < len(header_indices) else len(rows)
        sec_title = sheet_name
        for back_i in range(h_idx - 1, max(-1, h_idx - 4), -1):
            cell_vals = [str(c).strip() for c in rows[back_i] if c is not None and str(c).strip()]
            if cell_vals and len(cell_vals) <= 2:
                candidate_title = " ".join(cell_vals).strip()
                if (
                    len(candidate_title) >= 3
                    and not candidate_title.lower().startswith("total")
                    and not re.match(r"^[\d.,\sR\$]+$", candidate_title)
                ):
                    sec_title = candidate_title
                    break
        sec_rows = rows[h_idx:end_idx]
        sections.append((sec_title, sec_rows))

    return sections


def parse_b3_spreadsheet(file_bytes: bytes, filename: str = "planilha_b3.xlsx") -> Dict[str, Any]:
    """
    Parses any official B3 Investor Area spreadsheet (.xlsx or .csv):
      - Relatório de Posição (Ações, FIIs, ETFs, BDRs, Tesouro Direto, Renda Fixa)
      - Relatório de Negociação (Compras e Vendas)
      - Relatório de Movimentação / Proventos (Dividendos, Rendimentos, JCP)
    Returns structured items ready for preview or database sync.
    """
    sheets = load_workbook_sheets(file_bytes, filename)
    if not sheets:
        raise ValueError("Não foi possível ler nenhuma aba ou linha válida no arquivo enviado.")

    positions: List[Dict[str, Any]] = []
    transactions: List[Dict[str, Any]] = []
    fixed_income: List[Dict[str, Any]] = []
    dividends: List[Dict[str, Any]] = []
    detected_sheets: List[Dict[str, Any]] = []

    today_str = date.today().strftime("%Y-%m-%d")

    for sheet_raw_name, raw_rows in sheets.items():
        sub_sections = _partition_sheet_into_sections(sheet_raw_name, raw_rows)
        for sheet_name, rows in sub_sections:
            if len(rows) < 2:
                continue

            h_idx, headers = _find_header_row(rows)
            data_rows = rows[h_idx + 1:]
            sh_norm = _norm_text(sheet_name)
            header_joined = " | ".join(headers)

            sheet_items_count = 0
            sheet_kind = "DESCONHECIDO"

            # -----------------------------------------------------------------
            # 1. RELATÓRIO DE NEGOCIAÇÃO (Extrato de Compras e Vendas B3)
            # -----------------------------------------------------------------
            is_negociacao = (
                "negociacao" in sh_norm
                or "negocio" in sh_norm
                or "ordens" in sh_norm
                or "trade" in sh_norm
                or ("data do negocio" in header_joined and "tipo de movimentacao" in header_joined)
                or ("data do negocio" in header_joined and "compra" in header_joined)
                or ("tipo de operacao" in header_joined and "quantidade" in header_joined)
                or ("c/v" in header_joined and "quantidade" in header_joined)
            )
            if is_negociacao:
                sheet_kind = "NEGOCIACAO"
                c_date = _col_idx(headers, ["data do negocio", "data negocio", "data do pregao", "data pregao", "data operacao", "data"])
                c_type = _col_idx(headers, ["tipo de movimentacao", "tipo movimentacao", "tipo de operacao", "tipo operacao", "compra/venda", "c/v", "natureza", "operacao", "tipo"])
                c_code = _col_idx(headers, ["codigo de negociacao", "codigo negociacao", "codigo do ativo", "codigo", "ticker", "ativo"])
                c_prod = _col_idx(headers, ["produto", "especificacao", "especificacao do ativo", "nome", "empresa"])
                c_inst = _col_idx(headers, ["instituicao", "instituicao intermediadora", "corretora", "agente de custodia", "custodiante"])
                c_qty = _col_idx(headers, ["quantidade", "qtd", "cotas"])
                c_price = _col_idx(headers, ["preco unitario", "preco medio", "preco", "preco (r$)", "valor unitario", "taxa"])
                c_val = _col_idx(headers, ["valor total", "valor total (r$)", "valor da operacao", "total", "valor"])

                for r in data_rows:
                    raw_code = str(_get_cell(r, c_code, "")).strip()
                    raw_prod = str(_get_cell(r, c_prod, "")).strip()
                    if not raw_code and not raw_prod:
                        continue
                    if _norm_text(raw_code) in ("total", "subtotal") or _norm_text(raw_prod) in ("total", "subtotal"):
                        continue

                    ticker, name = _extract_ticker_and_name(raw_prod, raw_code)
                    if not ticker or len(ticker) < 2:
                        continue

                    qty = _parse_br_float(_get_cell(r, c_qty, 0))
                    if qty <= 0:
                        continue

                    unit_price = _parse_br_float(_get_cell(r, c_price, 0))
                    total_val = _parse_br_float(_get_cell(r, c_val, 0))
                    if unit_price <= 0 and total_val > 0 and qty > 0:
                        unit_price = round(total_val / qty, 4)
                    if unit_price <= 0:
                        continue

                    raw_type = _norm_text(_get_cell(r, c_type, "compra"))
                    tx_type = "SELL" if ("venda" in raw_type or raw_type in ("v", "sell", "debito", "saida")) else "BUY"
                    tx_date = _parse_br_date(_get_cell(r, c_date, today_str))
                    broker = _clean_broker_name(_get_cell(r, c_inst, "B3 / Corretora"))
                    asset_class, currency = _infer_asset_class(ticker, sheet_name, name)

                    transactions.append({
                        "tx_date": tx_date,
                        "ticker": ticker,
                        "name": name,
                        "asset_class": asset_class,
                        "tx_type": tx_type,
                        "quantity": round(qty, 6),
                        "unit_price": round(unit_price, 4),
                        "total_brl": round(qty * unit_price, 2),
                        "currency": currency,
                        "broker": broker,
                        "notes": f"Importado B3 ({sheet_name})",
                    })
                    sheet_items_count += 1

                detected_sheets.append({"sheet": sheet_name, "kind": sheet_kind, "items": sheet_items_count})
                continue

            # -----------------------------------------------------------------
            # 2. RELATÓRIO DE MOVIMENTAÇÃO / PROVENTOS (Dividendos, JCP, Rendimentos)
            # -----------------------------------------------------------------
            is_movimentacao = (
                "movimentacao" in sh_norm
                or "provento" in sh_norm
                or "dividendo" in sh_norm
                or ("entrada/saida" in header_joined and "movimentacao" in header_joined)
                or ("tipo de provento" in header_joined)
            )
            if is_movimentacao:
                sheet_kind = "MOVIMENTACAO_PROVENTOS"
                c_io = _col_idx(headers, ["entrada/saida", "natureza", "credito/debito"])
                c_date = _col_idx(headers, ["data", "data do pagamento", "pagamento"])
                c_mov = _col_idx(headers, ["movimentacao", "tipo de provento", "evento", "tipo", "historico"])
                c_prod = _col_idx(headers, ["produto", "ativo", "codigo de negociacao", "ticker"])
                c_code = _col_idx(headers, ["codigo de negociacao", "ticker", "codigo"])
                c_inst = _col_idx(headers, ["instituicao", "corretora"])
                c_qty = _col_idx(headers, ["quantidade", "cotas", "qtd"])
                c_price = _col_idx(headers, ["preco unitario", "valor por cota", "valor unitario", "preco"])
                c_val = _col_idx(headers, ["valor da operacao", "valor liquido", "valor total", "valor", "total"])

                for r in data_rows:
                    raw_prod = str(_get_cell(r, c_prod, "")).strip()
                    raw_code = str(_get_cell(r, c_code, "")).strip() if c_code != c_prod else ""
                    if not raw_prod and not raw_code:
                        continue
                    if _norm_text(raw_prod) in ("total", "subtotal"):
                        continue

                    mov_raw = str(_get_cell(r, c_mov, "Rendimento")).strip()
                    mov_norm = _norm_text(mov_raw)
                    io_norm = _norm_text(_get_cell(r, c_io, "credito"))

                    # Check if this row is a Dividend/Yield/JCP payment
                    is_div = any(w in mov_norm for w in (
                        "rendimento", "dividendo", "juros sobre capital", "jcp",
                        "provento", "restituicao de capital", "amortizacao"
                    ))
                    if is_div:
                        ticker, name = _extract_ticker_and_name(raw_prod, raw_code)
                        if not ticker or len(ticker) < 2:
                            continue
                        qty = _parse_br_float(_get_cell(r, c_qty, 0))
                        unit_val = _parse_br_float(_get_cell(r, c_price, 0))
                        total_val = _parse_br_float(_get_cell(r, c_val, 0))

                        if total_val <= 0 and qty > 0 and unit_val > 0:
                            total_val = round(qty * unit_val, 2)
                        if qty <= 0 and total_val > 0:
                            qty = 1.0
                            unit_val = total_val
                        if unit_val <= 0 and total_val > 0 and qty > 0:
                            unit_val = round(total_val / qty, 4)

                        if total_val <= 0:
                            continue

                        if "juros" in mov_norm or "jcp" in mov_norm:
                            div_label = "JCP"
                        elif "rendimento" in mov_norm:
                            div_label = "Rendimento FII"
                        else:
                            div_label = "Dividendo"

                        pay_date = _parse_br_date(_get_cell(r, c_date, today_str))
                        dividends.append({
                            "pay_date": pay_date,
                            "ticker": ticker,
                            "div_type": div_label,
                            "amount_per_share": round(unit_val, 4),
                            "shares_qty": round(qty, 4),
                            "total_amount": round(total_val, 2),
                            "currency": "BRL",
                        })
                        sheet_items_count += 1
                    elif any(w in mov_norm for w in ("compra", "venda")) and "liquidacao" not in mov_norm:
                        # Also support buy/sell rows inside Movimentacao if present
                        ticker, name = _extract_ticker_and_name(raw_prod, raw_code)
                        qty = _parse_br_float(_get_cell(r, c_qty, 0))
                        unit_val = _parse_br_float(_get_cell(r, c_price, 0))
                        total_val = _parse_br_float(_get_cell(r, c_val, 0))
                        if unit_val <= 0 and total_val > 0 and qty > 0:
                            unit_val = round(total_val / qty, 4)
                        if ticker and qty > 0 and unit_val > 0:
                            tx_type = "SELL" if ("venda" in mov_norm or "debito" in io_norm) else "BUY"
                            asset_class, currency = _infer_asset_class(ticker, sheet_name, name)
                            transactions.append({
                                "tx_date": _parse_br_date(_get_cell(r, c_date, today_str)),
                                "ticker": ticker,
                                "name": name,
                                "asset_class": asset_class,
                                "tx_type": tx_type,
                                "quantity": round(qty, 6),
                                "unit_price": round(unit_val, 4),
                                "total_brl": round(qty * unit_val, 2),
                                "currency": currency,
                                "broker": _clean_broker_name(_get_cell(r, c_inst, "B3 / Corretora")),
                                "notes": f"Importado B3 ({sheet_name})",
                            })
                            sheet_items_count += 1

                detected_sheets.append({"sheet": sheet_name, "kind": sheet_kind, "items": sheet_items_count})
                continue

            # -----------------------------------------------------------------
            # 3. TESOURO DIRETO & RENDA FIXA (CDB, LCI, LCA, CRI, CRA, Debêntures)
            # -----------------------------------------------------------------
            is_tesouro_or_rf = (
                "tesouro" in sh_norm
                or "renda fixa" in sh_norm
                or "cdb" in sh_norm
                or "lci" in sh_norm
                or "lca" in sh_norm
                or "cri" in sh_norm
                or "cra" in sh_norm
                or "debenture" in sh_norm
                or "indexador" in header_joined
                or ("data de vencimento" in header_joined and "codigo de negociacao" not in header_joined)
                or ("vencimento" in header_joined and "preco de fechamento" not in header_joined and "codigo de negociacao" not in header_joined)
            )
            if is_tesouro_or_rf:
                is_tesouro_sheet = "tesouro" in sh_norm
                sheet_kind = "TESOURO_DIRETO" if is_tesouro_sheet else "RENDA_FIXA"

                c_prod = _col_idx(headers, ["produto", "titulo", "nome", "ativo"])
                c_inst = _col_idx(headers, ["instituicao", "corretora", "custodiante"])
                c_iss = _col_idx(headers, ["emissor", "instituicao emissora"])
                c_idx = _col_idx(headers, ["indexador", "indice", "rentabilidade", "taxa"])
                c_rate = _col_idx(headers, ["taxa", "percentual", "rentabilidade contratada"])
                c_start = _col_idx(headers, ["data de emissao", "data de aplicacao", "data compra", "inicio", "data de inicio", "data da operacao", "data operacao", "data"])
                c_mat = _col_idx(headers, ["data de vencimento", "vencimento"])

                val_cols_updated = [
                    "valor atualizado curva", "valor atualizado", "valor bruto",
                    "valor atualizado fechamento", "valor atualizado mtm", "saldo", "valor total", "saldo bruto"
                ]
                val_cols_applied = ["valor aplicado", "valor investido", "aplicado", "custo total", "custo", "valor de compra", "valor de aplicacao"]
                c_net = _col_idx(headers, ["valor liquido"])

                for r in data_rows:
                    raw_prod = str(_get_cell(r, c_prod, "")).strip()
                    if not raw_prod or _norm_text(raw_prod) in ("total", "subtotal"):
                        continue

                    upd_val = 0.0
                    for cand in val_cols_updated:
                        idx = _col_idx(headers, [cand])
                        if idx >= 0:
                            v = _parse_br_float(_get_cell(r, idx, 0))
                            if v > 0:
                                upd_val = v
                                break

                    inv_val = 0.0
                    for cand in val_cols_applied:
                        idx = _col_idx(headers, [cand])
                        if idx >= 0:
                            v = _parse_br_float(_get_cell(r, idx, 0))
                            if v > 0:
                                inv_val = v
                                break

                    net_val = _parse_br_float(_get_cell(r, c_net, 0)) if c_net >= 0 else 0.0
                    if upd_val <= 0 and net_val > 0:
                        upd_val = net_val

                    # Valor de mercado (âncora) = valor atualizado; custo = valor aplicado (quando a B3 informa)
                    market_value = upd_val if upd_val > 0 else inv_val
                    if market_value <= 0:
                        continue
                    cost_amount = inv_val if inv_val > 0 else market_value

                    prod_norm = _norm_text(raw_prod)
                    raw_idx = _norm_text(_get_cell(r, c_idx, ""))
                    raw_rate = _parse_br_float(_get_cell(r, c_rate, 0))

                    # Determine asset_class (TESOURO vs CAIXINHA vs RENDA_FIXA)
                    if is_tesouro_sheet or "tesouro" in prod_norm or "ntn-b" in prod_norm or "lft" in prod_norm or "ltn" in prod_norm:
                        a_class = "TESOURO"
                    elif "caixinha" in prod_norm or "reserva" in prod_norm:
                        a_class = "CAIXINHA"
                    else:
                        a_class = "RENDA_FIXA"

                    # Determine indexer and rate
                    if "ipca" in raw_idx or "ipca" in prod_norm or "ntn-b" in prod_norm:
                        indexer = "IPCA"
                        rate_val = raw_rate if 0 < raw_rate < 30 else 6.25
                    elif "selic" in raw_idx or "selic" in prod_norm or "lft" in prod_norm:
                        indexer = "SELIC"
                        rate_val = raw_rate if 0 <= raw_rate < 10 else 0.10
                    elif "pre" in raw_idx or "prefixado" in prod_norm or "ltn" in prod_norm:
                        indexer = "PRE"
                        rate_val = raw_rate if 0 < raw_rate < 40 else 11.5
                    else:
                        indexer = "CDI"
                        rate_val = raw_rate if raw_rate >= 50 else 100.0

                    # Extract rate from string if embedded (e.g., "110% CDI" or "IPCA + 6,50%")
                    m_pct = re.search(r"(\d+(?:[.,]\d+)?)\s*%", f"{raw_prod} {_get_cell(r, c_idx, '')}")
                    if m_pct and raw_rate <= 0:
                        parsed_pct = _parse_br_float(m_pct.group(1))
                        if parsed_pct > 0:
                            rate_val = parsed_pct

                    broker = _clean_broker_name(_get_cell(r, c_inst, "Tesouro / Banco"))
                    issuer = str(_get_cell(r, c_iss, "")).strip() or ("Tesouro Nacional" if a_class == "TESOURO" else broker)
                    mat_raw = _get_cell(r, c_mat, "2028-12-31")
                    mat_dt = _parse_br_date(mat_raw) if str(mat_raw).strip() else "2028-12-31"

                    tax_exempt = 1 if any(k in prod_norm for k in ("lci", "lca", "cri", "cra", "incentivada")) else 0

                    # Data de início real (emissão/aplicação) quando a B3 informa; usada para a faixa de IR.
                    raw_start = str(_get_cell(r, c_start, "")).strip() if c_start >= 0 else ""
                    has_start_date = False
                    start_dt = today_str
                    if raw_start and raw_start not in ("", "-"):
                        parsed_start = _parse_br_date(raw_start)
                        if parsed_start:
                            start_dt = parsed_start
                            has_start_date = True

                    has_cost_amount = (inv_val > 0)

                    fixed_income.append({
                        "title_name": raw_prod[:80],
                        "issuer": issuer[:50],
                        "asset_class": a_class,
                        "indexer": indexer,
                        "rate_value": round(rate_val, 2),
                        "invested_amount": round(cost_amount, 2),
                        "market_value": round(market_value, 2),
                        "market_net_value": round(net_val, 2) if net_val > 0 else None,
                        "market_value_date": today_str,
                        "start_date": start_dt,
                        "has_start_date": has_start_date,
                        "has_cost_amount": has_cost_amount,
                        "maturity_date": mat_dt,
                        "daily_liquidity": 1 if a_class in ("CAIXINHA", "TESOURO") else 0,
                        "tax_exempt": tax_exempt,
                        "broker": broker,
                    })
                    sheet_items_count += 1

                detected_sheets.append({"sheet": sheet_name, "kind": sheet_kind, "items": sheet_items_count})
                continue

            # -----------------------------------------------------------------
            # 4. POSIÇÃO EM RENDA VARIÁVEL (Ações, FIIs, ETFs, BDRs)
            # -----------------------------------------------------------------
            c_code = _col_idx(headers, ["codigo de negociacao", "codigo", "ticker", "ativo"])
            c_prod = _col_idx(headers, ["produto", "nome", "empresa", "fundo", "especificacao"])
            c_inst = _col_idx(headers, ["instituicao", "corretora", "custodiante"])
            c_cnpj = _col_idx(headers, ["cnpj da empresa", "cnpj do fundo", "cnpj"])
            c_tipo = _col_idx(headers, ["tipo", "classe", "categoria", "segmento"])
            c_qty = _col_idx(headers, ["quantidade", "quantidade disponivel", "qtd", "cotas"])
            c_avg_price = _col_idx(headers, ["preco medio", "preco medio de compra", "preco de compra", "custo medio", "pm"])
            c_price = _col_idx(headers, ["preco de fechamento", "cotacao", "ultimo preco", "preco unitario", "preco"])
            c_val = _col_idx(headers, ["valor atualizado", "valor total", "saldo", "valor"])

            if c_qty >= 0 and (c_code >= 0 or c_prod >= 0):
                sheet_kind = "POSICAO_RENDA_VARIAVEL"
                for r in data_rows:
                    raw_code = str(_get_cell(r, c_code, "")).strip()
                    raw_prod = str(_get_cell(r, c_prod, "")).strip()
                    if not raw_code and not raw_prod:
                        continue
                    if _norm_text(raw_code) in ("total", "subtotal") or _norm_text(raw_prod) in ("total", "subtotal"):
                        continue

                    ticker, name = _extract_ticker_and_name(raw_prod, raw_code)
                    if not ticker or len(ticker) < 2:
                        continue

                    qty = _parse_br_float(_get_cell(r, c_qty, 0))
                    if qty <= 0:
                        continue

                    avg_price = _parse_br_float(_get_cell(r, c_avg_price, 0)) if c_avg_price >= 0 else 0.0
                    close_price = _parse_br_float(_get_cell(r, c_price, 0)) if c_price >= 0 else 0.0
                    total_val = _parse_br_float(_get_cell(r, c_val, 0))

                    # Se informou Preço Médio, ele é o custo contábil. Caso contrário usa fechamento ou total/qtd
                    unit_price = avg_price if avg_price > 0 else (close_price if close_price > 0 else (total_val / qty if total_val > 0 and qty > 0 else 0.0))

                    raw_tipo = str(_get_cell(r, c_tipo, "")).strip()
                    asset_class, currency = _infer_asset_class(ticker, sheet_name, name, raw_tipo)
                    broker = _clean_broker_name(_get_cell(r, c_inst, "B3 / Corretora"))
                    cnpj = str(_get_cell(r, c_cnpj, "")).strip()

                    positions.append({
                        "ticker": ticker,
                        "name": name,
                        "asset_class": asset_class,
                        "currency": currency,
                        "quantity": round(qty, 6),
                        "unit_price": round(unit_price, 4),
                        "avg_price": round(avg_price, 4) if avg_price > 0 else 0.0,
                        "has_explicit_avg_price": avg_price > 0,
                        "close_price": round(close_price, 4),
                        "total_brl": round(total_val if total_val > 0 else qty * (close_price or unit_price), 2),
                        "broker": broker,
                        "cnpj": cnpj,
                        "sheet": sheet_name,
                    })
                    sheet_items_count += 1

                detected_sheets.append({"sheet": sheet_name, "kind": sheet_kind, "items": sheet_items_count})

    # Consolidate positions that appear across multiple rows/accounts in the same Position file
    consolidated_positions_map: Dict[str, Dict[str, Any]] = {}
    for p in positions:
        key = p["ticker"]
        if key not in consolidated_positions_map:
            consolidated_positions_map[key] = dict(p)
        else:
            existing = consolidated_positions_map[key]
            new_qty = existing["quantity"] + p["quantity"]
            new_total = existing["total_brl"] + p["total_brl"]
            existing["quantity"] = round(new_qty, 6)
            existing["total_brl"] = round(new_total, 2)
            if new_qty > 0:
                if existing.get("avg_price", 0) > 0 and p.get("avg_price", 0) > 0:
                    existing["avg_price"] = round((existing["avg_price"] * (new_qty - p["quantity"]) + p["avg_price"] * p["quantity"]) / new_qty, 4)
                elif p.get("avg_price", 0) > 0:
                    existing["avg_price"] = p["avg_price"]
                existing["has_explicit_avg_price"] = existing.get("avg_price", 0) > 0
                existing["unit_price"] = existing.get("avg_price") or round(new_total / new_qty, 4)
            if p["broker"] not in existing["broker"]:
                existing["broker"] = f"{existing['broker']} / {p['broker']}"[:45]

    consolidated_positions = list(consolidated_positions_map.values())

    # Annotate positions with existing portfolio holding status in DB
    sold_positions: List[Dict[str, Any]] = []
    try:
        conn = get_conn()
        existing_holdings = get_current_portfolio_holdings(conn)
        catalog_rows = conn.execute("SELECT ticker, name, asset_class FROM assets_catalog").fetchall()
        conn.close()
        cat_map = {r["ticker"].upper().strip(): dict(r) for r in catalog_rows}

        for cp in consolidated_positions:
            h = existing_holdings.get(cp["ticker"].upper().strip(), {})
            cp["existing_qty"] = round(h.get("quantity", 0.0), 6)
            cp["existing_pm"] = round(h.get("avg_price", 0.0), 2)
            cp["has_existing"] = h.get("quantity", 0.0) > 1e-6

        fname = str(filename or "").lower()
        detected_kinds = {s.get("kind") for s in detected_sheets if s.get("kind")}
        detected_sheets_names = {_norm_text(s.get("sheet")) for s in detected_sheets if s.get("sheet")}
        is_pos_filename = any(k in fname for k in ("posicao", "posiçao", "custodia", "carteira"))
        has_pos_kind = "POSICAO_RENDA_VARIAVEL" in detected_kinds
        has_pos_sheet_names = any(any(k in s for k in ("posicao", "posiçao", "custodia", "acoes", "fundo", "fii", "renda variavel")) for s in detected_sheets_names)
        is_position_report = bool(consolidated_positions or has_pos_kind or is_pos_filename or has_pos_sheet_names)

        if is_position_report:
            if len(detected_sheets_names) == 1 and "acoes" in list(detected_sheets_names)[0] and not is_pos_filename:
                reconciled_classes = {"ACAO_BR", "ETF_BR"}
            elif len(detected_sheets_names) == 1 and any(k in list(detected_sheets_names)[0] for k in ("fundo", "fii")) and not is_pos_filename:
                reconciled_classes = {"FII"}
            else:
                reconciled_classes = set(B3_CUSTODY_TRANSACTION_CLASSES)

            imported_keys = {cp["ticker"].upper().strip() for cp in consolidated_positions}
            for t, h in existing_holdings.items():
                q = float(h.get("quantity", 0.0))
                if q > 1e-6 and t not in imported_keys:
                    c_info = cat_map.get(t, {})
                    cls = (c_info.get("asset_class") or h.get("asset_class") or "").upper().strip()
                    if cls in reconciled_classes and cls in B3_CUSTODY_TRANSACTION_CLASSES:
                        sold_positions.append({
                            "ticker": t,
                            "name": c_info.get("name") or t,
                            "asset_class": cls,
                            "quantity": round(q, 6),
                            "avg_price": round(h.get("avg_price", 0.0), 2),
                            "broker": h.get("broker", "B3 Custódia"),
                        })
    except Exception:
        pass

    total_value_brl = (
        sum(p["total_brl"] for p in consolidated_positions)
        + sum(t["total_brl"] for t in transactions if t["tx_type"] == "BUY")
        + sum(f.get("market_value") or f["invested_amount"] for f in fixed_income)
    )

    return {
        "filename": filename,
        "sheets_detected": detected_sheets,
        "detected_sheets": detected_sheets,
        "counts": {
            "positions": len(consolidated_positions),
            "transactions": len(transactions),
            "fixed_income": len(fixed_income),
            "dividends": len(dividends),
            "sold_positions": len(sold_positions),
            "total_items": len(consolidated_positions) + len(transactions) + len(fixed_income) + len(dividends),
        },
        "estimated_value_brl": round(total_value_brl, 2),
        "positions": consolidated_positions,
        "sold_positions": sold_positions,
        "transactions": transactions,
        "fixed_income": fixed_income,
        "dividends": dividends,
    }


def _ensure_asset_in_catalog(cur, ticker: str, name: str, asset_class: str, currency: str, unit_price: float, cnpj: str = ""):
    exists = cur.execute("SELECT ticker, cnpj, last_price FROM assets_catalog WHERE ticker = ?", (ticker,)).fetchone()
    if exists:
        if cnpj and not exists["cnpj"]:
            cur.execute("UPDATE assets_catalog SET cnpj = ? WHERE ticker = ?", (cnpj, ticker))
        if unit_price > 0 and (not exists["last_price"] or float(exists["last_price"] or 0) <= 0):
            cur.execute("UPDATE assets_catalog SET last_price = ? WHERE ticker = ?", (unit_price, ticker))
        return

    sector = "Fundo Imobiliário" if asset_class == "FII" else ("Ações B3" if asset_class == "ACAO_BR" else "Mercado")
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    cur.execute("""
        INSERT OR IGNORE INTO assets_catalog (
            ticker, name, asset_class, currency, sector, cnpj, last_price, change_pct,
            volume_day, dy_12m, pl, pvp, last_dividend_per_share, updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        ticker,
        name or ticker,
        asset_class,
        currency,
        sector,
        cnpj or "",
        unit_price if unit_price > 0 else 10.0,
        0.0,
        1000000,
        0.0,
        0.0,
        1.0,
        0.0,
        now_str,
    ))


def apply_b3_sync(parsed_data: Dict[str, Any], sync_mode: str = "MERGE", positions_avg_prices: Optional[Dict[str, float]] = None, user_id: int = 1) -> Dict[str, Any]:
    """
    Applies parsed B3 spreadsheet items to the database.
    sync_mode:
      - 'MERGE': Smart deduplication & position reconciliation:
        * If real trades from Negociação are imported, they populate historical buys/sells and calculate real PM.
        * If a custody position is imported and already exists in SIM with a known Preço Médio,
          it PRESERVES the Preço Médio 100% and only adjusts the custody quantity.
        * Any prior synthetic placeholder transactions for tickers in Negociação are safely replaced.
      - 'REPLACE': Clears existing portfolio positions for the imported categories and replaces with the B3 file.
    """
    ensure_b3_sync_tables()
    mode = (sync_mode or "MERGE").upper().strip()
    positions_avg_prices = positions_avg_prices or {}
    macro = fetch_macro_indicators()
    usd_fx = float(macro.get("usd_brl", 5.72))
    today_str = date.today().strftime("%Y-%m-%d")

    positions = parsed_data.get("positions", [])
    transactions = parsed_data.get("transactions", [])
    fixed_income = parsed_data.get("fixed_income", [])
    dividends = parsed_data.get("dividends", [])

    conn = get_conn()
    cur = conn.cursor()

    pos_synced = 0
    tx_synced = 0
    fi_synced = 0
    div_synced = 0
    skipped = 0

    if mode == "REPLACE":
        # CRÍTICO: No modo REPLACE, apenas classes ESTRITAMENTE custodiadas na B3 (ACAO_BR, FII, ETF_BR)
        # que estejam efetivamente presentes no arquivo importado devem ser apagadas.
        # NUNCA apagar classes fora da B3: STOCK_US (ações EUA), CRIPTO, CAIXINHA ou lançamentos manuais não-B3.
        imported_classes = set()
        if positions:
            imported_classes.update(p.get("asset_class") for p in positions if p.get("asset_class"))
        if transactions:
            imported_classes.update(t.get("asset_class") for t in transactions if t.get("asset_class"))

        # Interseção estrita: só apaga classes de renda variável/fundos negociadas e custodiadas na B3
        classes_to_clear = imported_classes.intersection(B3_CUSTODY_TRANSACTION_CLASSES)
        if classes_to_clear:
            placeholders = ",".join("?" for _ in classes_to_clear)
            cur.execute(f"DELETE FROM transactions WHERE user_id = ? AND asset_class IN ({placeholders})", (user_id, *classes_to_clear))

        if fixed_income:
            fi_classes = set(f.get("asset_class") for f in fixed_income if f.get("asset_class"))
            fi_to_clear = fi_classes.intersection(B3_CUSTODY_FIXED_INCOME_CLASSES)
            if fi_to_clear:
                fi_placeholders = ",".join("?" for _ in fi_to_clear)
                cur.execute(f"DELETE FROM fixed_income WHERE user_id = ? AND asset_class IN ({fi_placeholders})", (user_id, *fi_to_clear))

        if dividends:
            # Apaga apenas proventos em BRL dos tickers presentes no relatório importado
            div_tickers = {d["ticker"].upper() for d in dividends if d.get("ticker")}
            if div_tickers:
                div_placeholders = ",".join("?" for _ in div_tickers)
                cur.execute(
                    f"DELETE FROM dividends_received WHERE user_id = ? AND currency = 'BRL' AND UPPER(ticker) IN ({div_placeholders})",
                    (user_id, *div_tickers),
                )

    # -------------------------------------------------------------------------
    # 1. PROCESS INDIVIDUAL B3 TRANSACTIONS (from B3 Negociação Report) FIRST
    # -------------------------------------------------------------------------
    # Negociação contains real historical buys & sells with exact trade execution prices.
    # When importing Negociação for a ticker, any prior synthetic placeholder transactions
    # (created when only Posição was available) are cleaned up so the real history takes over.
    if transactions:
        tx_tickers = {tx["ticker"].upper() for tx in transactions}
        for t in tx_tickers:
            cur.execute("""
                DELETE FROM transactions
                WHERE user_id = ? AND UPPER(ticker) = ? AND (
                    notes LIKE '%Sincroniza%o Posi%o B3%'
                    OR notes LIKE '%Posi%o B3%'
                    OR notes LIKE '%Ajuste de quantidade B3%'
                    OR notes LIKE '%Ajuste de cust%dia B3%'
                    OR notes LIKE '%Ajuste%'
                    OR notes LIKE '%Pre%o M%dio%'
                )
            """, (user_id, t))

        for tx in transactions:
            ticker = tx["ticker"].upper()
            tx_date = tx["tx_date"]
            tx_type = tx["tx_type"]
            qty = float(tx["quantity"])
            unit_price = float(tx["unit_price"])
            asset_class = tx["asset_class"]
            currency = tx["currency"]
            fx = usd_fx if currency == "USD" else 1.0

            _ensure_asset_in_catalog(cur, ticker, tx.get("name", ticker), asset_class, currency, unit_price)

            if mode == "MERGE":
                dup = cur.execute("""
                    SELECT id FROM transactions
                    WHERE user_id = ? AND tx_date = ? AND UPPER(ticker) = ? AND tx_type = ?
                      AND ABS(quantity - ?) < 0.0001 AND ABS(unit_price - ?) < 0.01
                """, (user_id, tx_date, ticker, tx_type, qty, unit_price)).fetchone()
                if dup:
                    skipped += 1
                    continue

            cur.execute("""
                INSERT INTO transactions (user_id, tx_date, ticker, asset_class, tx_type, quantity, unit_price, fees, currency, fx_rate, broker, notes)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                user_id,
                tx_date,
                ticker,
                asset_class,
                tx_type,
                qty,
                unit_price,
                0.0,
                currency,
                fx,
                tx.get("broker", "B3 / Corretora"),
                tx.get("notes", "Sincronização Negociação B3"),
            ))
            tx_synced += 1

    # -------------------------------------------------------------------------
    # 2. PROCESS CONSOLIDATED B3 POSITIONS & RECONCILIATION
    # -------------------------------------------------------------------------
    imported_tickers = set()

    if positions:
        # Re-evaluate current portfolio state in DB (takes into account newly applied Negociação trades!)
        existing_state = get_current_portfolio_holdings(cur, user_id=user_id)

        for p in positions:
            ticker = p["ticker"].upper().strip()
            imported_tickers.add(ticker)
            target_qty = float(p["quantity"])
            close_price = float(p.get("close_price") or p.get("unit_price") or 0.0)
            asset_class = p["asset_class"]
            currency = p["currency"]
            fx = usd_fx if currency == "USD" else 1.0

            _ensure_asset_in_catalog(cur, ticker, p["name"], asset_class, currency, close_price or 10.0, p.get("cnpj", ""))

            curr_holding = existing_state.get(ticker, {})
            curr_qty = max(0.0, curr_holding.get("quantity", 0.0))
            existing_pm = curr_holding.get("avg_price", 0.0)

            diff = round(target_qty - curr_qty, 6)

            if abs(diff) <= 1e-6:
                # Quantity already matches B3 custody perfectly; keep existing PM untouched!
                skipped += 1
                continue

            # CRUCIAL: If the asset ALREADY has an existing Preço Médio in SIM:
            # PRESERVE THE PREÇO MÉDIO 100% AND ONLY ADJUST THE CUSTODY QUANTITY!
            if existing_pm > 0:
                if diff > 0:
                    # Quantity increased in custody: buy the difference at existing_pm so new PM == existing_pm!
                    cur.execute("""
                        INSERT INTO transactions (user_id, tx_date, ticker, asset_class, tx_type, quantity, unit_price, fees, currency, fx_rate, broker, notes)
                        VALUES (?, ?, ?, ?, 'BUY', ?, ?, 0.0, ?, ?, ?, ?)
                    """, (
                        user_id,
                        today_str,
                        ticker,
                        asset_class,
                        diff,
                        round(existing_pm, 4),
                        currency,
                        fx,
                        p.get("broker", "B3 Custódia"),
                        f"Ajuste de quantidade B3 (Preço Médio mantido em R$ {existing_pm:.2f})",
                    ))
                    pos_synced += 1
                else:
                    # Quantity decreased in custody (diff < 0): sell the difference at existing_pm
                    cur.execute("""
                        INSERT INTO transactions (user_id, tx_date, ticker, asset_class, tx_type, quantity, unit_price, fees, currency, fx_rate, broker, notes)
                        VALUES (?, ?, ?, ?, 'SELL', ?, ?, 0.0, ?, ?, ?, ?)
                    """, (
                        user_id,
                        today_str,
                        ticker,
                        asset_class,
                        abs(diff),
                        round(existing_pm, 4),
                        currency,
                        fx,
                        p.get("broker", "B3 Custódia"),
                        "Ajuste de quantidade B3",
                    ))
                    pos_synced += 1
            else:
                # Ticker is NEW to SIM (no prior trades/PM in DB):
                # Check if the user specified a custom Preço Médio via preview or spreadsheet
                user_pm = float(positions_avg_prices.get(ticker, 0.0)) if positions_avg_prices else 0.0
                sheet_pm = float(p.get("avg_price", 0.0))

                if user_pm > 0:
                    chosen_pm = user_pm
                    pm_note = f"Sincronização Posição B3 (Preço Médio informado: R$ {chosen_pm:.2f})"
                elif sheet_pm > 0:
                    chosen_pm = sheet_pm
                    pm_note = f"Sincronização Posição B3 (Preço Médio da planilha: R$ {chosen_pm:.2f})"
                else:
                    # Fallback to closing price when no PM is available
                    chosen_pm = close_price if close_price > 0 else 10.0
                    pm_note = "Posição B3 (Preço Médio estimado pelo fechamento - importe Negociação para valor real)"

                tx_type = "BUY" if diff > 0 else "SELL"
                cur.execute("""
                    INSERT INTO transactions (user_id, tx_date, ticker, asset_class, tx_type, quantity, unit_price, fees, currency, fx_rate, broker, notes)
                    VALUES (?, ?, ?, ?, ?, ?, ?, 0.0, ?, ?, ?, ?)
                """, (
                    user_id,
                    today_str,
                    ticker,
                    asset_class,
                    tx_type,
                    abs(diff),
                    round(chosen_pm, 4),
                    currency,
                    fx,
                    p.get("broker", "B3 Custódia"),
                    pm_note,
                ))
                pos_synced += 1

    # -------------------------------------------------------------------------
    # RECONCILIAÇÃO B3: Ativos vendidos / que não constam mais na custódia B3
    # -------------------------------------------------------------------------
    zeroed_tickers: List[str] = []
    if mode == "MERGE":
        fname = str(parsed_data.get("filename") or "").lower()
        sheets_det = parsed_data.get("sheets_detected", []) or parsed_data.get("detected_sheets", [])
        detected_kinds = {s.get("kind") for s in sheets_det if s.get("kind")}
        detected_sheets_names = {_norm_text(s.get("sheet")) for s in sheets_det if s.get("sheet")}
        is_pos_filename = any(k in fname for k in ("posicao", "posiçao", "custodia", "carteira"))
        has_pos_kind = "POSICAO_RENDA_VARIAVEL" in detected_kinds
        has_pos_sheet_names = any(any(k in s for k in ("posicao", "posiçao", "custodia", "acoes", "fundo", "fii", "renda variavel")) for s in detected_sheets_names)
        is_position_report = bool(positions or has_pos_kind or is_pos_filename or has_pos_sheet_names)

        if is_position_report:
            if len(detected_sheets_names) == 1 and "acoes" in list(detected_sheets_names)[0] and not is_pos_filename:
                reconciled_classes = {"ACAO_BR", "ETF_BR"}
            elif len(detected_sheets_names) == 1 and any(k in list(detected_sheets_names)[0] for k in ("fundo", "fii")) and not is_pos_filename:
                reconciled_classes = {"FII"}
            else:
                reconciled_classes = set(B3_CUSTODY_TRANSACTION_CLASSES)

            # Re-read existing portfolio holdings (accounts for any newly inserted position trades)
            current_holdings = get_current_portfolio_holdings(cur, user_id=user_id)
            for t, h_info in current_holdings.items():
                curr_q = float(h_info.get("quantity", 0.0))
                if curr_q <= 1e-6:
                    continue
                if t in imported_tickers:
                    continue

                h_class = (h_info.get("asset_class") or "").upper().strip()
                cat_row = cur.execute("SELECT asset_class, last_price FROM assets_catalog WHERE UPPER(ticker) = ?", (t,)).fetchone()
                cat_class = (cat_row["asset_class"] if cat_row and cat_row["asset_class"] else "").upper().strip()

                target_class = cat_class if cat_class in B3_CUSTODY_TRANSACTION_CLASSES else (h_class if h_class in B3_CUSTODY_TRANSACTION_CLASSES else "")

                if target_class and target_class in reconciled_classes:
                    existing_pm = float(h_info.get("avg_price", 0.0))
                    sell_p = existing_pm if existing_pm > 0 else (float(cat_row["last_price"]) if cat_row and cat_row["last_price"] else 10.0)
                    broker = h_info.get("broker") or "B3 Custódia"
                    cur.execute("""
                        INSERT INTO transactions (user_id, tx_date, ticker, asset_class, tx_type, quantity, unit_price, fees, currency, fx_rate, broker, notes)
                        VALUES (?, ?, ?, ?, 'SELL', ?, ?, 0.0, 'BRL', 1.0, ?, 'Ajuste: ativo não consta mais na custódia B3')
                    """, (user_id, today_str, t, target_class, curr_q, round(sell_p, 4), broker))
                    pos_synced += 1
                    zeroed_tickers.append(t)

    # -------------------------------------------------------------------------
    # 3. PROCESS FIXED INCOME & TESOURO DIRETO
    # -------------------------------------------------------------------------
    for fi in fixed_income:
        market_value = fi.get("market_value")
        market_net_value = fi.get("market_net_value")
        market_value_date = fi.get("market_value_date") or today_str
        invested_amt = fi["invested_amount"]

        if mode == "MERGE":
            fi_broker = (fi.get("broker") or "").strip()
            dup = None
            if fi_broker:
                dup = cur.execute("""
                    SELECT id, invested_amount, start_date, rate_value, broker FROM fixed_income
                    WHERE user_id = ? AND LOWER(title_name) = LOWER(?) AND LOWER(issuer) = LOWER(?)
                      AND maturity_date = ? AND LOWER(broker) = LOWER(?)
                """, (user_id, fi["title_name"], fi["issuer"], fi["maturity_date"], fi_broker)).fetchone()
                if not dup:
                    dup = cur.execute("""
                        SELECT id, invested_amount, start_date, rate_value, broker FROM fixed_income
                        WHERE user_id = ? AND LOWER(title_name) = LOWER(?)
                          AND maturity_date = ? AND LOWER(broker) = LOWER(?)
                    """, (user_id, fi["title_name"], fi["maturity_date"], fi_broker)).fetchone()
            else:
                dup = cur.execute("""
                    SELECT id, invested_amount, start_date, rate_value, broker FROM fixed_income
                    WHERE user_id = ? AND LOWER(title_name) = LOWER(?) AND LOWER(issuer) = LOWER(?)
                      AND maturity_date = ? AND (broker IS NULL OR broker = '')
                """, (user_id, fi["title_name"], fi["issuer"], fi["maturity_date"])).fetchone()

            if dup:
                cost_to_save = invested_amt
                # Preserve existing invested_amount if the imported one was just a fallback equal to market_value
                if not fi.get("has_cost_amount", True) and dup["invested_amount"] > 0:
                    cost_to_save = dup["invested_amount"]
                elif dup["invested_amount"] > 0 and (cost_to_save <= 0 or abs(cost_to_save - (market_value or 0)) < 0.01):
                    cost_to_save = dup["invested_amount"]

                # Preserve existing start_date if incoming had no explicit date from sheet
                start_to_save = fi["start_date"]
                if not fi.get("has_start_date", True) and dup["start_date"]:
                    start_to_save = dup["start_date"]

                cur.execute("""
                    UPDATE fixed_income
                    SET invested_amount = ?, start_date = ?, rate_value = ?, broker = ?,
                        market_value = ?, market_net_value = ?, market_value_date = ?
                    WHERE id = ? AND user_id = ?
                """, (
                    cost_to_save, start_to_save, fi["rate_value"], fi["broker"],
                    market_value, market_net_value, market_value_date, dup["id"], user_id,
                ))
                fi_synced += 1
                continue

        cur.execute("""
            INSERT INTO fixed_income (
                user_id, title_name, issuer, asset_class, indexer, rate_value,
                invested_amount, start_date, maturity_date, daily_liquidity, tax_exempt, broker,
                market_value, market_net_value, market_value_date
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            user_id,
            fi["title_name"],
            fi["issuer"],
            fi["asset_class"],
            fi["indexer"],
            fi["rate_value"],
            invested_amt,
            fi["start_date"],
            fi["maturity_date"],
            fi["daily_liquidity"],
            fi["tax_exempt"],
            fi["broker"],
            market_value,
            market_net_value,
            market_value_date if market_value else "",
        ))
        fi_synced += 1

    # 4. Process Dividends / Proventos
    for d in dividends:
        if mode == "MERGE":
            dup = cur.execute("""
                SELECT id FROM dividends_received
                WHERE user_id = ? AND pay_date = ? AND ticker = ? AND ABS(total_amount - ?) < 0.02
            """, (user_id, d["pay_date"], d["ticker"], d["total_amount"])).fetchone()
            if dup:
                skipped += 1
                continue

        cur.execute("""
            INSERT INTO dividends_received (user_id, pay_date, ticker, div_type, amount_per_share, shares_qty, total_amount, currency)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            user_id,
            d["pay_date"],
            d["ticker"],
            d["div_type"],
            d["amount_per_share"],
            d["shares_qty"],
            d["total_amount"],
            d["currency"],
        ))
        div_synced += 1

    sheets_for_log = parsed_data.get("sheets_detected") or parsed_data.get("detected_sheets") or []
    report_kinds = ", ".join(s["sheet"] for s in sheets_for_log if s.get("sheet"))
    summary_msg = (
        f"Sincronização concluída ({mode}): {pos_synced} posições de custódia, "
        f"{tx_synced} negociações, {fi_synced} títulos de renda fixa/Tesouro e {div_synced} proventos sincronizados."
    )
    if zeroed_tickers:
        summary_msg += f" ({len(zeroed_tickers)} posição(ões) vendida(s) zerada(s): {', '.join(zeroed_tickers)})."
    if skipped > 0:
        summary_msg += f" ({skipped} registros idênticos ignorados sem duplicar)."

    cur.execute("""
        INSERT INTO b3_sync_logs (
            user_id, synced_at, filename, sync_mode, report_types,
            positions_synced, transactions_synced, fixed_income_synced, dividends_synced,
            skipped_duplicates, summary_text
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        user_id,
        datetime.now().strftime("%d/%m/%Y %H:%M:%S"),
        parsed_data.get("filename", "planilha_b3.xlsx"),
        mode,
        report_kinds,
        pos_synced,
        tx_synced,
        fi_synced,
        div_synced,
        skipped,
        summary_msg,
    ))

    conn.commit()
    conn.close()

    # Trigger background quote update safely without holding any DB locks
    try:
        from .market_service import sync_yfinance_quotes_background
        sync_yfinance_quotes_background()
    except Exception:
        pass

    return {
        "status": "ok",
        "sync_mode": mode,
        "positions_synced": pos_synced,
        "transactions_synced": tx_synced,
        "fixed_income_synced": fi_synced,
        "dividends_synced": div_synced,
        "skipped_duplicates": skipped,
        "zeroed_positions": zeroed_tickers,
        "summary": summary_msg,
    }


def get_b3_sync_logs(limit: int = 10, user_id: int = 1) -> List[Dict[str, Any]]:
    ensure_b3_sync_tables()
    conn = get_conn()
    rows = conn.execute("SELECT * FROM b3_sync_logs WHERE user_id = ? ORDER BY id DESC LIMIT ?", (user_id, limit)).fetchall()
    conn.close()
    return [dict(r) for r in rows]


def _style_header_row(ws, headers: List[str]):
    ws.append(headers)
    header_fill = PatternFill(start_color="111317", end_color="111317", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="3ECF8E")
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="left", vertical="center")


def _auto_fit_columns(ws):
    for col in ws.columns:
        max_len = 12
        col_letter = col[0].column_letter
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = min(48, len(val_str) + 4)
        ws.column_dimensions[col_letter].width = max_len


def generate_b3_template_xlsx() -> bytes:
    """
    Generates an official B3-formatted Excel template (.xlsx) with sample tabs:
    - Acoes
    - Fundo de Investimento
    - Tesouro Direto
    - Renda Fixa
    - Movimentacao
    - Negociacao
    """
    wb = openpyxl.Workbook()
    # 1. Acoes
    ws_acoes = wb.active
    ws_acoes.title = "Acoes"
    _style_header_row(ws_acoes, [
        "Produto", "Instituição", "Código de Negociação", "CNPJ da Empresa",
        "Tipo", "Quantidade", "Preço de Fechamento", "Valor Atualizado"
    ])
    ws_acoes.append(["PETR4 - PETROLEO BRASILEIRO S.A. PETROBRAS", "XP INVESTIMENTOS CCTVM S/A", "PETR4", "33.000.167/0001-01", "PN", 100, 36.85, 3685.00])
    ws_acoes.append(["WEGE3 - WEG S.A.", "NU INVEST CORRETORA DE VALORES S.A.", "WEGE3", "84.429.695/0001-11", "ON", 50, 52.90, 2645.00])
    ws_acoes.append(["BBAS3 - BANCO DO BRASIL S.A.", "BTG PACTUAL CTVM S/A", "BBAS3", "00.000.000/0001-91", "ON", 120, 28.45, 3414.00])
    _auto_fit_columns(ws_acoes)

    # 2. Fundo de Investimento (FIIs & ETFs)
    ws_fii = wb.create_sheet(title="Fundo de Investimento")
    _style_header_row(ws_fii, [
        "Produto", "Instituição", "Código de Negociação", "CNPJ do Fundo",
        "Tipo", "Quantidade", "Preço de Fechamento", "Valor Atualizado"
    ])
    ws_fii.append(["MXRF11 - FII MAXI RENDA", "XP INVESTIMENTOS CCTVM S/A", "MXRF11", "97.521.225/0001-25", "Cotas", 300, 9.96, 2988.00])
    ws_fii.append(["HGLG11 - CSHG LOGISTICA FII", "BTG PACTUAL CTVM S/A", "HGLG11", "11.728.688/0001-47", "Cotas", 20, 158.40, 3168.00])
    ws_fii.append(["IVVB11 - ISHARES S&P 500 FDO INV COTAS FDO INDICE", "NU INVEST CORRETORA DE VALORES S.A.", "IVVB11", "19.909.560/0001-91", "ETF", 10, 358.00, 3580.00])
    _auto_fit_columns(ws_fii)

    # 3. Tesouro Direto
    ws_td = wb.create_sheet(title="Tesouro Direto")
    _style_header_row(ws_td, [
        "Produto", "Instituição", "Indexador", "Data de Vencimento",
        "Quantidade", "Valor Aplicado", "Valor Bruto", "Valor Líquido"
    ])
    ws_td.append(["Tesouro Selic 2029", "NU INVEST CORRETORA DE VALORES S.A.", "SELIC", "01/03/2029", 1.5, 15000.00, 16420.00, 16100.00])
    ws_td.append(["Tesouro IPCA+ 2035", "XP INVESTIMENTOS CCTVM S/A", "IPCA", "15/05/2035", 2.0, 8500.00, 9240.00, 9080.00])
    _auto_fit_columns(ws_td)

    # 4. Renda Fixa
    ws_rf = wb.create_sheet(title="Renda Fixa")
    _style_header_row(ws_rf, [
        "Produto", "Instituição", "Emissor", "Indexador", "Taxa",
        "Data de Emissão", "Data de Vencimento", "Valor Aplicado", "Valor Atualizado"
    ])
    ws_rf.append(["Caixinha Reserva 100% CDI", "Nubank", "Nubank", "CDI", 100.0, "10/01/2026", "31/12/2028", 10000.00, 10850.00])
    ws_rf.append(["LCI Banco Inter 94% CDI", "Banco Inter", "Banco Inter", "CDI", 94.0, "15/03/2026", "15/03/2028", 5000.00, 5390.00])
    _auto_fit_columns(ws_rf)

    # 5. Movimentacao (Proventos)
    ws_mov = wb.create_sheet(title="Movimentacao")
    _style_header_row(ws_mov, [
        "Entrada/Saída", "Data", "Movimentação", "Produto", "Instituição",
        "Quantidade", "Preço unitário", "Valor da Operação"
    ])
    ws_mov.append(["Credito", "15/10/2026", "Rendimento", "MXRF11 - FII MAXI RENDA", "XP INVESTIMENTOS CCTVM S/A", 300, 0.10, 30.00])
    ws_mov.append(["Credito", "15/10/2026", "Rendimento", "HGLG11 - CSHG LOGISTICA FII", "BTG PACTUAL CTVM S/A", 20, 1.10, 22.00])
    ws_mov.append(["Credito", "20/09/2026", "Dividendo", "PETR4 - PETROLEO BRASILEIRO S.A.", "XP INVESTIMENTOS CCTVM S/A", 100, 0.92, 92.00])
    _auto_fit_columns(ws_mov)

    # 6. Negociacao (Extrato de Compras e Vendas B3 - Histórico Real)
    ws_neg = wb.create_sheet(title="Negociacao")
    _style_header_row(ws_neg, [
        "Data do Negócio", "Tipo de Movimentação", "Mercado", "Prazo/Vencimento",
        "Instituição", "Código de Negociação", "Quantidade", "Preço (R$)", "Valor Total (R$)"
    ])
    ws_neg.append(["15/01/2026", "Compra", "Mercado à Vista", "", "XP INVESTIMENTOS CCTVM S/A", "PETR4", 100, 34.20, 3420.00])
    ws_neg.append(["10/02/2026", "Compra", "Mercado à Vista", "", "XP INVESTIMENTOS CCTVM S/A", "MXRF11", 114, 10.15, 1157.10])
    ws_neg.append(["05/03/2026", "Compra", "Mercado à Vista", "", "NU INVEST CORRETORA DE VALORES S.A.", "ITUB4", 4, 32.50, 130.00])
    ws_neg.append(["12/03/2026", "Compra", "Mercado à Vista", "", "NU INVEST CORRETORA DE VALORES S.A.", "PRIO3", 2, 65.00, 130.00])
    ws_neg.append(["20/03/2026", "Compra", "Mercado à Vista", "", "BANCO COOPERATIVO SICREDI S/A", "GARE11", 105, 9.10, 955.50])
    ws_neg.append(["25/03/2026", "Compra", "Mercado à Vista", "", "BANCO COOPERATIVO SICREDI S/A", "GGRC11", 14, 11.20, 156.80])
    ws_neg.append(["28/03/2026", "Compra", "Mercado à Vista", "", "BANCO COOPERATIVO SICREDI S/A", "BBAS3", 1, 26.50, 26.50])
    _auto_fit_columns(ws_neg)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def export_portfolio_b3_xlsx(portfolio_summary: Dict[str, Any]) -> bytes:
    """
    Exports the user's current SIM portfolio into a multi-sheet .xlsx workbook formatted like B3's Investor Area,
    so it can be backed up, edited in Excel, or shared and re-imported seamlessly.
    """
    wb = openpyxl.Workbook()
    positions = portfolio_summary.get("positions", [])
    fi_list = portfolio_summary.get("fixed_income", [])
    divs = portfolio_summary.get("dividends", [])
    txs = portfolio_summary.get("transactions", [])

    # 1. Aba Acoes (e Stocks/Cripto)
    ws_acoes = wb.active
    ws_acoes.title = "Acoes"
    _style_header_row(ws_acoes, [
        "Produto", "Instituição", "Código de Negociação", "Tipo",
        "Quantidade", "Preço Médio", "Preço de Fechamento", "Valor Atualizado"
    ])
    for p in positions:
        if p["asset_class"] in ("ACAO_BR", "STOCK_US", "CRIPTO"):
            ws_acoes.append([
                f"{p['ticker']} - {p['name']}",
                p.get("broker", "B3 / Corretora"),
                p["ticker"],
                p["asset_class"],
                p["quantity"],
                p["avg_price"],
                p["current_price"],
                p["current_total_brl"],
            ])
    _auto_fit_columns(ws_acoes)

    # 2. Aba Fundo de Investimento (FIIs & ETFs)
    ws_fii = wb.create_sheet(title="Fundo de Investimento")
    _style_header_row(ws_fii, [
        "Produto", "Instituição", "Código de Negociação", "Tipo",
        "Quantidade", "Preço Médio", "Preço de Fechamento", "Valor Atualizado"
    ])
    for p in positions:
        if p["asset_class"] in ("FII", "ETF_BR", "FUNDO_INV"):
            ws_fii.append([
                f"{p['ticker']} - {p['name']}",
                p.get("broker", "B3 / Corretora"),
                p["ticker"],
                "ETF" if p["asset_class"] == "ETF_BR" else "Cotas",
                p["quantity"],
                p["avg_price"],
                p["current_price"],
                p["current_total_brl"],
            ])
    _auto_fit_columns(ws_fii)

    # 3. Aba Tesouro Direto
    ws_td = wb.create_sheet(title="Tesouro Direto")
    _style_header_row(ws_td, [
        "Produto", "Instituição", "Indexador", "Taxa",
        "Data de Emissão", "Data de Vencimento", "Valor Aplicado", "Valor Líquido"
    ])
    for f in fi_list:
        if f.get("asset_class") == "TESOURO":
            ws_td.append([
                f["title_name"],
                f.get("broker", "Tesouro Nacional"),
                f["indexer"],
                f["rate_value"],
                f["start_date"],
                f["maturity_date"],
                f["invested_amount"],
                f.get("net_amount", f["invested_amount"]),
            ])
    _auto_fit_columns(ws_td)

    # 4. Aba Renda Fixa
    ws_rf = wb.create_sheet(title="Renda Fixa")
    _style_header_row(ws_rf, [
        "Produto", "Instituição", "Emissor", "Indexador", "Taxa",
        "Data de Emissão", "Data de Vencimento", "Valor Aplicado", "Valor Atualizado"
    ])
    for f in fi_list:
        if f.get("asset_class") != "TESOURO":
            ws_rf.append([
                f["title_name"],
                f.get("broker", f["issuer"]),
                f["issuer"],
                f["indexer"],
                f["rate_value"],
                f["start_date"],
                f["maturity_date"],
                f["invested_amount"],
                f.get("net_amount", f["invested_amount"]),
            ])
    _auto_fit_columns(ws_rf)

    # 5. Aba Movimentacao (Proventos)
    ws_mov = wb.create_sheet(title="Movimentacao")
    _style_header_row(ws_mov, [
        "Entrada/Saída", "Data", "Movimentação", "Produto",
        "Quantidade", "Preço unitário", "Valor da Operação"
    ])
    for d in divs:
        ws_mov.append([
            "Credito",
            d["pay_date"],
            d.get("div_type", "Rendimento"),
            d["ticker"],
            d["shares_qty"],
            d["amount_per_share"],
            d["total_amount"],
        ])
    _auto_fit_columns(ws_mov)

    # 6. Aba Negociacao (Histórico de Ordens)
    ws_neg = wb.create_sheet(title="Negociacao")
    _style_header_row(ws_neg, [
        "Data do Negócio", "Tipo de Movimentação", "Código de Negociação",
        "Instituição", "Quantidade", "Preço", "Valor"
    ])
    for tx in txs:
        ws_neg.append([
            tx["tx_date"],
            "Compra" if tx["tx_type"] == "BUY" else "Venda",
            tx["ticker"],
            tx.get("broker", "B3 / Corretora"),
            tx["quantity"],
            tx["unit_price"],
            round(float(tx["quantity"]) * float(tx["unit_price"]), 2),
        ])
    _auto_fit_columns(ws_neg)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
